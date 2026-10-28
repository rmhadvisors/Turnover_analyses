"""Excel / PDF exports and the month-wise report data."""

import io
import re
from datetime import datetime
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from app.services import report_export
from app.services.calculations import round2
from app.services.comparison_service import MonthlyComparison, MonthlySeries
from app.utils.indian_format import CRORE, LAKH
from tests.test_api import enter_worked_example, make_client, upload


def download(api, path, **params):
    response = api.get(path, params=params)
    assert response.status_code == 200, response.text
    return response


def test_excel_report_has_real_numbers_and_formats(api) -> None:
    client_id = make_client(api, "Acme Traders")
    enter_worked_example(api, client_id)
    response = download(api, f"/reports/export/{client_id}", fy="2025-26", unit="lakhs")
    assert "spreadsheetml" in response.headers["content-type"]
    assert (
        "Turnover_Comparison_Acme_Traders_FY2025-26.xlsx" in response.headers["content-disposition"]
    )

    sheet = load_workbook(io.BytesIO(response.content))["Comparison"]
    assert sheet["A1"].value == "TURNOVER COMPARISON – FY 2025-26"
    assert sheet["A2"].value == "Client: Acme Traders"
    assert [c.value for c in sheet[5]][:7] == [
        "Particular", "Unit", "FY 2024-25", "FY 2025-26", "Difference", "Change %", "Status",
    ]  # fmt: skip

    turnover = {c.column_letter: c for c in sheet[6]}
    assert turnover["B"].value == "₹ L"  # every figure in the row is in that unit
    assert turnover["C"].value == 80 and turnover["D"].value == 100  # numbers, in lakhs
    assert turnover["E"].value == 20
    assert isinstance(turnover["C"].value, (int, float))  # not text
    assert "₹" in turnover["C"].number_format and '" L' in turnover["C"].number_format
    assert "↑" in turnover["E"].number_format and "↓" in turnover["E"].number_format
    assert turnover["F"].value == 0.25 and turnover["F"].number_format == "0.00%"
    assert turnover["G"].value == "Significant Increase"

    net_profit = {c.column_letter: c for c in sheet[9]}
    assert net_profit["A"].value == "Net Profit"
    assert net_profit["E"].value == -2 and abs(net_profit["F"].value - (-0.2857)) < 0.00005
    assert net_profit["F"].font.color.rgb.endswith("B42318")  # red: down
    assert net_profit["G"].value == "Significant Decrease"

    purchases = {c.column_letter: c for c in sheet[7]}
    assert purchases["G"].value == "New / No comparison" and purchases["F"].value == "—"

    text = " ".join(str(c.value) for row in sheet.iter_rows() for c in row if c.value)
    assert "GST registration - goods" in text  # limit alert listed below the table


def test_excel_units_and_full_indian_format(api) -> None:
    client_id = make_client(api)
    enter_worked_example(api, client_id)

    def turnover_row(unit):
        content = download(api, f"/reports/export/{client_id}", fy="2025-26", unit=unit).content
        return {c.column_letter: c for c in load_workbook(io.BytesIO(content))["Comparison"][6]}

    crores = turnover_row("crores")
    assert crores["B"].value == "₹ Cr"
    assert crores["D"].value == 1 and '" Cr' in crores["D"].number_format
    full = turnover_row("full")
    assert full["B"].value == "₹" and full["D"].value == 10000000
    assert "##\\,##\\,##" in full["D"].number_format  # Indian grouping format, still numeric
    auto = turnover_row("auto")
    # 1 crore switches the whole row to Cr, so the Unit cell holds for both years
    assert auto["B"].value == "₹ Cr" and auto["C"].value == 0.8 and auto["D"].value == 1


def test_export_is_built_fresh_and_stamped_in_ist(api, sample_dir) -> None:
    client_id = make_client(api, "Sharma Traders")
    upload(api, sample_dir / "sharma_traders_sales_register.xlsx", client_id, "sales_register")
    before = datetime.now(report_export.IST).replace(second=0, microsecond=0)
    response = download(api, f"/reports/export/{client_id}", fy="2025-26")
    after = datetime.now(report_export.IST)
    assert response.headers["cache-control"] == "no-store"
    workbook = load_workbook(io.BytesIO(response.content))
    stamps = {workbook["Comparison"]["A3"].value, workbook["Month-wise"]["A3"].value}
    assert len(stamps) == 1  # one build time on every sheet
    (stamp,) = stamps
    built = datetime.strptime(stamp, "Generated %d-%b-%Y %H:%M IST").replace(
        tzinfo=report_export.IST
    )
    assert before <= built <= after  # the time of this request, in India time
    assert workbook["Comparison"]["A4"].value == "Amounts in lakhs / crores."
    pdf = download(api, f"/reports/export/{client_id}", fy="2025-26", format="pdf")
    assert pdf.headers["cache-control"] == "no-store"


def test_export_rejects_bad_parameters(api) -> None:
    client_id = make_client(api)
    assert api.get(f"/reports/export/{client_id}", params={"fy": "bad"}).status_code == 422
    good = {"fy": "2025-26"}
    assert (
        api.get(f"/reports/export/{client_id}", params={**good, "unit": "yards"}).status_code == 422
    )
    assert (
        api.get(f"/reports/export/{client_id}", params={**good, "format": "doc"}).status_code == 422
    )
    assert api.get("/reports/export/999", params=good).status_code == 404


def test_pdf_report_is_a_valid_document(api, sample_dir) -> None:
    client_id = make_client(api, "Sharma Traders")
    upload(api, sample_dir / "sharma_traders_sales_register.xlsx", client_id, "sales_register")
    upload(
        api, sample_dir / "sharma_traders_purchase_register.xlsx", client_id, "purchase_register"
    )
    response = download(api, f"/reports/export/{client_id}", fy="2025-26", format="pdf")
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")
    assert response.content.rstrip().endswith(b"%%EOF")
    assert len(re.findall(rb"/Type\s*/Page\b", response.content)) >= 1
    assert b"DejaVu" in response.content  # embedded font that has the rupee sign


def test_pdf_and_excel_work_without_any_vouchers(api) -> None:
    client_id = make_client(api)
    enter_worked_example(api, client_id)  # manual entry only: no month-wise data
    assert (
        download(api, f"/reports/export/{client_id}", fy="2025-26", format="pdf").content[:4]
        == b"%PDF"
    )
    content = download(api, f"/reports/export/{client_id}", fy="2025-26").content
    assert load_workbook(io.BytesIO(content)).sheetnames == ["Comparison"]


def test_monthly_data_and_workbook_chart(api, sample_dir) -> None:
    client_id = make_client(api, "Sharma Traders")
    upload(api, sample_dir / "sharma_traders_sales_register.xlsx", client_id, "sales_register")
    upload(
        api, sample_dir / "sharma_traders_purchase_register.xlsx", client_id, "purchase_register"
    )

    monthly = download(api, f"/reports/monthly/{client_id}", fy="2025-26").json()
    assert monthly["has_data"] is True
    assert monthly["months"] == [
        "Apr",
        "May",
        "Jun",
        "Jul",
        "Aug",
        "Sep",
        "Oct",
        "Nov",
        "Dec",
        "Jan",
        "Feb",
        "Mar",
    ]
    assert monthly["current"]["fy"] == "2025-26" and monthly["previous"]["fy"] == "2024-25"
    # month-wise totals add up to the FY turnover / purchases held in the report
    figures = {f["fy"]: f for f in api.get(f"/entries/{client_id}").json()}
    assert sum(Decimal(v) for v in monthly["current"]["sales"]) == Decimal(
        figures["2025-26"]["turnover"]
    )
    assert sum(Decimal(v) for v in monthly["previous"]["purchases"]) == Decimal(
        figures["2024-25"]["purchases"]
    )

    content = download(api, f"/reports/export/{client_id}", fy="2025-26").content
    workbook = load_workbook(io.BytesIO(content))
    assert workbook.sheetnames == ["Comparison", "Month-wise"]
    assert len(workbook["Month-wise"]._charts) == 2  # current and previous FY
    assert isinstance(workbook["Month-wise"]["B5"].value, float)


UNIT_SCALE = {"₹": Decimal(1), "₹ L": LAKH, "₹ Cr": CRORE}


def rupees(value: float, scale: Decimal) -> Decimal:
    """A stored (scaled) cell value back in rupees, to the paisa."""
    return round2(Decimal(repr(value)) * scale)


def month_wise_totals_match_comparison(workbook) -> dict[str, tuple[Decimal, Decimal]]:
    """Assert the Month-wise Total row equals the Comparison sheet's annual figures."""
    comparison, monthly = workbook["Comparison"], workbook["Month-wise"]
    assert monthly["A17"].value == "Total"
    totals = {}
    for r, label, cols in ((6, "Turnover", "BCDE"), (7, "Purchases", "FGHI")):
        assert comparison.cell(row=r, column=1).value == label
        scale = UNIT_SCALE[comparison.cell(row=r, column=2).value]
        annual = tuple(rupees(comparison.cell(row=r, column=c).value, scale) for c in (3, 4))
        totals[label] = tuple(rupees(monthly[f"{c}17"].value, LAKH) for c in cols[:2])
        assert totals[label] == annual, label
        difference = rupees(comparison.cell(row=r, column=5).value, scale)
        assert rupees(monthly[f"{cols[2]}17"].value, LAKH) == difference, label
        assert monthly[f"{cols[3]}17"].value == comparison.cell(row=r, column=6).value, label
    return totals


@pytest.mark.parametrize(
    ("name", "slug", "ext"),
    [
        ("Sharma Traders", "sharma_traders", "xlsx"),
        ("Patel Engineering Works", "patel_engineering_works", "csv"),
        ("Iyer Consulting LLP", "iyer_consulting_llp", "xlsx"),
    ],
)
@pytest.mark.parametrize("unit", ["auto", "lakhs", "crores", "full"])
def test_month_wise_totals_equal_comparison_sheet(api, sample_dir, name, slug, ext, unit) -> None:
    client_id = make_client(api, name)
    for kind in ("sales", "purchase"):
        upload(api, sample_dir / f"{slug}_{kind}_register.{ext}", client_id, f"{kind}_register")
    content = download(api, f"/reports/export/{client_id}", fy="2025-26", unit=unit).content
    totals = month_wise_totals_match_comparison(load_workbook(io.BytesIO(content)))
    figures = {f["fy"]: f for f in api.get(f"/entries/{client_id}").json()}
    for i, fy in enumerate(("2024-25", "2025-26")):
        assert totals["Turnover"][i] == Decimal(figures[fy]["turnover"])
        assert totals["Purchases"][i] == Decimal(figures[fy]["purchases"])


def test_month_wise_layout_change_pct_and_block_rule() -> None:
    def series(fy, sales, purchases):
        return MonthlySeries(fy, [Decimal(v) for v in sales], [Decimal(v) for v in purchases])

    zeros = ["0"] * 12
    monthly = MonthlyComparison(
        client_id=1, client_name="X", months=("Apr", "May", *("M",) * 10), has_data=True,
        previous=series("2024-25", ["100000", "0", *zeros[2:]], ["200000", "50000", *zeros[2:]]),
        current=series("2025-26", ["150000", "70000", *zeros[2:]], ["100000", "50000", *zeros[2:]]),
    )  # fmt: skip
    workbook = Workbook()
    report_export._monthly_sheet(workbook, monthly, "2025-26", "2024-25", "Generated now")
    ws = workbook["Month-wise"]
    assert [c.value for c in ws[4]] == [
        "Month",
        "Sales FY 2024-25", "Sales FY 2025-26", "Sales Difference", "Sales Change %",
        "Purchases FY 2024-25", "Purchases FY 2025-26", "Purchases Difference", "Purchases Change %",
    ]  # fmt: skip
    # Difference in lakhs, with the Comparison sheet's arrow format, green up / red down
    assert (
        ws["D5"].value == 0.5 and "↑" in ws["D5"].number_format and '" L' in ws["D5"].number_format
    )
    assert ws["D5"].font.color.rgb.endswith("137333")
    assert ws["H5"].value == -1 and ws["H5"].font.color.rgb.endswith("B42318")
    # Change %
    assert ws["E5"].value == 0.5 and ws["E5"].number_format == "0.00%"
    assert ws["E5"].font.color.rgb.endswith("137333")  # green: up
    assert ws["I5"].value == -0.5 and ws["I5"].font.color.rgb.endswith("B42318")  # red: down
    assert ws["D6"].value == 0.7 and ws["E6"].value == "—"  # no baseline: amount, but no %
    assert ws["H6"].value == 0 and ws["H6"].font.color is None  # unchanged: no colour
    assert ws["I6"].value == 0 and ws["I6"].font.color.type != "rgb"
    assert ws["D7"].value == 0 and ws["E7"].value == "—"  # zero both years
    # Total row holds values (not formulas), its own Difference and Change %
    assert [ws[f"{c}17"].value for c in "BCDE"] == [1, 2.2, 1.2, 1.2]
    assert [ws[f"{c}17"].value for c in "FGHI"] == [2.5, 1.5, -1, -0.4]
    assert ws["A17"].font.b and ws["C17"].font.b and ws["D17"].font.b and ws["E17"].font.b
    for r in range(4, 18):  # rule between the sales and purchases blocks
        assert ws.cell(row=r, column=5).border.right.style == "thin"
        assert ws.cell(row=r, column=5).border.right.color.rgb.endswith("1F3A5F")


def test_monthly_without_vouchers_is_empty(api) -> None:
    client_id = make_client(api)
    enter_worked_example(api, client_id)
    monthly = download(api, f"/reports/monthly/{client_id}", fy="2025-26").json()
    assert monthly["has_data"] is False
    assert set(monthly["current"]["sales"]) == {"0"}


def test_summary_export_sorted_by_biggest_mover(api, sample_dir) -> None:
    ids = {}
    for name, slug, ext in (
        ("Sharma Traders", "sharma_traders", "xlsx"),
        ("Patel Engineering Works", "patel_engineering_works", "csv"),
        ("Iyer Consulting LLP", "iyer_consulting_llp", "xlsx"),
    ):
        ids[name] = make_client(api, name)
        upload(api, sample_dir / f"{slug}_sales_register.{ext}", ids[name], "sales_register")
    make_client(api, "No Data Yet")

    content = download(api, "/reports/summary/export", fy="2025-26", unit="lakhs").content
    sheet = load_workbook(io.BytesIO(content))["Summary"]
    assert sheet.auto_filter.ref == "A4:L8"  # incl. the TDS columns
    names = [sheet.cell(row=r, column=1).value for r in range(5, 9)]
    assert names == [
        "Patel Engineering Works",
        "Sharma Traders",
        "Iyer Consulting LLP",
        "No Data Yet",
    ]
    patel = {c.column_letter: c for c in sheet[5]}
    assert patel["D"].number_format == "0.00%" and patel["D"].value < -0.3
    assert patel["E"].value == "Significant Decrease"
    assert isinstance(patel["C"].value, float)
    assert sheet["D8"].value == "—"  # no comparison sorts last
