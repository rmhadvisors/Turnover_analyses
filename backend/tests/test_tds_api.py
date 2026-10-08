"""TDS REST API: rate master, mapping approval, report, alert detail, actions, exports."""

import io
from decimal import Decimal as D

from openpyxl import load_workbook

from tests.test_tds_service import FILES

FY = "2025-26"


def _client(api) -> int:
    client_id = api.post("/clients", json={"name": "Synthetic TDS Firm"}).json()["id"]
    api.put(f"/clients/{client_id}/profile", json={"gstin": "27AAAFZ9999Z1Z5"})
    api.post("/entries", json={
        "client_id": client_id, "previous_fy": "2023-24", "current_fy": "2024-25",
        "previous_turnover": "100000000", "current_turnover": "150000000",
    })  # fmt: skip
    files = [("files", (name, content)) for name, content in FILES]
    response = api.post(
        "/imports/tally-json/preview", data={"client_id": str(client_id)}, files=files
    )
    assert response.status_code == 200, response.text
    token = response.json()["preview_token"]
    confirm = api.post(
        "/imports/tally-json/confirm", data={"client_id": str(client_id), "preview_token": token}
    )
    assert confirm.status_code == 200, confirm.text
    return client_id


def _approve(api, client_id: int) -> None:
    response = api.post(f"/tds/clients/{client_id}/mapping/approve", json={"approved_by": "CA"})
    assert response.status_code == 200 and response.json()["approved_by"] == "CA"


def test_rate_master_crud_with_effective_dates(api) -> None:
    body = api.get("/tds/sections").json()
    assert "Finance Act 2025" in body["note"] and "s.393" in body["note"]
    rows = body["sections"]
    c194 = next(r for r in rows if r["key"] == "194C" and r["effective_from"] == "2025-04-01")
    assert D(c194["single_threshold"]) == 30000 and D(c194["rate_other"]) == 2

    new = {**c194, "effective_from": "2026-04-01", "rate_other": "3"}
    created = api.post("/tds/sections", json=new)
    assert created.status_code == 201
    assert api.post("/tds/sections", json=new).status_code == 409  # same key + date
    edited = api.put(f"/tds/sections/{created.json()['id']}", json={**new, "rate_other": "2.5"})
    assert D(edited.json()["rate_other"]) == D("2.5")
    assert api.delete(f"/tds/sections/{created.json()['id']}").status_code == 204
    bad = api.post("/tds/sections", json={**new, "key": "194Z"})
    assert bad.status_code == 422


def test_mapping_review_and_approval(api) -> None:
    client_id = _client(api)
    mapping = api.get(f"/tds/clients/{client_id}/mapping").json()
    assert mapping["approved_at"] is None
    names = {r["name"]: r for r in mapping["rows"]}
    assert names["Mystery Expenses"]["role"] == "unmapped"
    changed = api.put(f"/tds/clients/{client_id}/mapping", json={"changes": [
        {"match_type": "ledger", "name": "Mystery Expenses", "role": "excluded"},
    ]})  # fmt: skip
    assert {r["name"]: r for r in changed.json()["rows"]}["Mystery Expenses"]["source"] == "user"
    missing_section = api.put(f"/tds/clients/{client_id}/mapping", json={"changes": [
        {"match_type": "ledger", "name": "Mystery Expenses", "role": "base"},
    ]})  # fmt: skip
    assert missing_section.status_code == 422
    assert api.get("/alerts", params={"kind": "tds"}).json() == []  # nothing before approval
    _approve(api, client_id)
    assert api.get("/alerts", params={"kind": "tds"}).json() != []


def test_report_reproduces_the_checker_layout_and_totals(api) -> None:
    client_id = _client(api)
    report = api.get(f"/tds/clients/{client_id}/report", params={"fy": FY}).json()
    assert report["payer"]["s194q"]["applies"] is True
    sections = {s["key"]: s for s in report["sections"]}
    rows_194c = {r["party"]: r for r in sections["194C"]["rows"]}
    silent = rows_194c["Silent Contractor"]
    assert (silent["pan_available"], silent["crossed"], silent["rate"]) == (
        "No",
        "Yes - deduct TDS",
        "20%",
    )
    assert D(silent["tds"]) == 7000 and D(silent["shortfall"]) == 7000
    # subtotal and grand total of the shortfall
    sub = sections["194C"]["subtotal"]
    assert D(sub["shortfall"]) == D(600) + D(7000) + D(
        1200
    )  # short + silent + truck (1% of 1.2 lakh)
    assert D(report["totals"]["shortfall"]) == D(sub["shortfall"])
    assert D(report["totals"]["tds"]) == D(1200) + D(1200) + D(7000) + D(3000) + D(1200)

    for fmt, magic in (("xlsx", b"PK"), ("pdf", b"%PDF")):
        response = api.get(
            f"/tds/clients/{client_id}/report/export", params={"fy": FY, "format": fmt}
        )
        assert response.status_code == 200 and response.content[:4].startswith(magic)
    sheet = load_workbook(io.BytesIO(api.get(
        f"/tds/clients/{client_id}/report/export", params={"fy": FY, "format": "xlsx"}).content))["TDS Applicability"]  # fmt: skip
    assert sheet["A6"].value == "Party / Payee name"


def test_alert_detail_panel_and_actions(api) -> None:
    client_id = _client(api)
    _approve(api, client_id)
    alerts = api.get("/alerts", params={"kind": "tds", "severity": "critical"}).json()
    silent = next(a for a in alerts if a["party"] == "Silent Contractor")
    assert (silent["section"], silent["severity"], silent["metric_label"]) == (
        "194C",
        "critical",
        "TDS 194C: Silent Contractor",
    )
    assert D(silent["shortfall"]) == 7000
    q_alerts = api.get("/alerts", params={"kind": "tds", "section": "194Q"}).json()
    assert [(a["new_status"], a["severity"]) for a in q_alerts] == [("tds_ok", "info")]
    short = api.get("/alerts", params={"kind": "tds", "severity": "high"}).json()
    assert [a["party"] for a in short] == ["Short Contractor"]

    detail = api.get(f"/tds/alerts/{silent['id']}/detail", params={"as_of": "2025-10-06"}).json()
    assert detail["party"]["pan"] is None and detail["party"]["pan_source"] == "not available"
    assert detail["section"]["why"]["applies"] is True
    assert detail["threshold"]["crossing_test"] == "single payment"
    assert detail["threshold"]["crossed_voucher_no"] == "C3"
    steps = {s["step"]: s for s in detail["calculation"]["steps"]}
    assert "206AA" in steps["Rate applied"]["note"] and D(steps["Shortfall"]["value"]) == 7000
    assert [v["is_crossing"] for v in detail["vouchers"]] == [True]
    items = {i["item"]: i for i in detail["compliance"]["items"]}
    assert "Not tax advice" in detail["compliance"]["disclaimer"]
    assert items["Deposit due date"]["figure"] == "2025-06-07"
    assert (
        D(items["Interest - late / no deduction (s.201(1A)(i), 1% per month or part)"]["figure"])
        == 420
    )
    assert D(items["Disallowance u/s 40(a)(ia)"]["figure"]) == 10500

    for fmt in ("xlsx", "pdf"):
        response = api.get(f"/tds/clients/{client_id}/detail/export", params={
            "fy": FY, "party_key": "guid-silent", "section": "194C", "format": fmt})  # fmt: skip
        assert response.status_code == 200

    # mark TDS deducted outside Tally -> the status changes to correctly deducted
    marked = api.put(f"/tds/clients/{client_id}/resolution", json={
        "fy": FY, "party_key": "guid-silent", "section_key": "194C", "note": "Deducted in a later voucher", "marked_by": "CA"})  # fmt: skip
    assert marked.status_code == 200
    latest = api.get("/alerts", params={"kind": "tds", "severity": "info"}).json()
    assert any(a["party"] == "Silent Contractor" and a["new_status"] == "tds_ok" for a in latest)

    ack = api.post(
        f"/tds/alerts/{silent['id']}/acknowledge",
        json={"acknowledged_by": "CA", "note": "Spoke to client"},
    )
    assert ack.json()["acknowledged"] and ack.json()["note"] == "Spoke to client"
    counts = api.get("/tds/alert-counts", params={"fy": FY}).json()
    assert set(counts) == {"critical", "high", "warning", "info"}


def test_party_facts_and_payer_overrides(api) -> None:
    client_id = _client(api)
    bad = api.put(
        f"/tds/clients/{client_id}/party", json={"party_key": "guid-silent", "pan": "123"}
    )
    assert bad.status_code == 422
    api.put(
        f"/tds/clients/{client_id}/party", json={"party_key": "guid-silent", "pan": "abcfs1234k"}
    )
    detail = api.get(
        f"/tds/clients/{client_id}/detail",
        params={"fy": FY, "party_key": "guid-silent", "section": "194C"},
    ).json()
    assert (detail["party"]["pan"], detail["party"]["pan_source"]) == (
        "ABCFS1234K",
        "entered by user",
    )
    assert detail["calculation"]["rate"] == "2%"  # firm with PAN now

    payer = api.get(f"/tds/clients/{client_id}/payer", params={"fy": FY}).json()
    assert payer["constitution"] == "Firm" and payer["s194q"]["applies"] is True
    saved = api.put(
        f"/tds/clients/{client_id}/payer",
        json={"fy": FY, "constitution": "Individual", "audit_liable": False},
    )
    assert saved.json()["others"]["applies"] is False
    report = api.get(f"/tds/clients/{client_id}/report", params={"fy": FY}).json()
    statuses = {r["status"] for s in report["sections"] if s["key"] == "194C" for r in s["rows"]}
    assert statuses == {"tds_not_applicable"}


def test_client_report_and_summary_carry_the_tds_block(api) -> None:
    client_id = _client(api)
    report = api.get(f"/reports/client/{client_id}", params={"fy": FY}).json()
    assert report["tds"]["parties_crossed"] == 5 and D(report["tds"]["not_deducted"]) == D(8800)
    row = next(
        r
        for r in api.get("/reports/summary", params={"fy": FY}).json()
        if r["client_id"] == client_id
    )
    assert row["tds_parties_crossed"] == 5 and D(row["tds_not_deducted"]) == D(8800)
    content = api.get(f"/reports/export/{client_id}", params={"fy": FY, "format": "xlsx"}).content
    assert "TDS" in load_workbook(io.BytesIO(content)).sheetnames
    assert (
        api.get(f"/reports/export/{client_id}", params={"fy": FY, "format": "pdf"}).status_code
        == 200
    )


def test_settings_approaching_pct(api) -> None:
    assert D(api.get("/tds/settings").json()["approaching_pct"]) == 80
    assert (
        D(api.put("/tds/settings", json={"approaching_pct": "90"}).json()["approaching_pct"]) == 90
    )
    assert api.put("/tds/settings", json={"approaching_pct": "0"}).status_code == 422
