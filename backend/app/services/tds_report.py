"""'TDS Applicability' report (the Applicability Checker sheet, one row per party per
section, with what was actually deducted) and the Excel / PDF exports of the report and of
one party's working."""

from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.services import tds_compliance as tc
from app.services import tds_detail
from app.services import tds_engine as te
from app.services.report_export import (
    _BOX,
    FMT_INDIAN_FULL,
    _header_row,
    _pdf_footer,
    _register_fonts,
    _styles,
)
from app.services.tds_service import ClientTds, PartySection

SEVERITY_FILL = {
    "tds_unidentified": "FDEBD9",
    "tds_not_deducted": "FDE7E7",
    "tds_short_deducted": "FDEBD9",
    "tds_approaching": "FFF1CC",
    "tds_ok": "E3F4E7",
    "tds_excess": "DCE8FB",
    "tds_nil": "E3F4E7",
}


def _counts(row: PartySection) -> bool:
    """Rows that enter the totals: an identified party, section applicable (or undetermined)."""
    return row.party.identified and row.decision.applies is not False


def report_row(row: PartySection) -> dict:
    ev = row.evaluation
    return {
        "party_key": row.party.key,
        "party": row.party.name,
        "pan": row.party.pan,
        "pan_source": row.party.pan_source,
        "section": row.section_key,
        "nature": f"{row.section_key} - {ev.rule.nature}",
        "payee_type": te.PAYEE_LABELS.get(row.party.payee_type, row.party.payee_type),
        "pan_available": "Yes" if row.party.pan_available else "No",
        "single_payment": str(ev.largest_single),
        "aggregate": str(ev.aggregate),
        "threshold": te.threshold_text(ev.rule),
        "crossed": "Yes - deduct TDS" if ev.crossed else "No - below threshold",
        "crossing_test": ev.crossing_test.value if ev.crossing_test else None,
        "crossed_on": ev.crossed_on.isoformat() if ev.crossed_on else None,
        "crossed_voucher_no": ev.crossed_voucher_no,
        "base": str(ev.base),
        "rate": tds_detail.rate_text(row),
        "tds": str(ev.tds),
        "deducted": str(row.deducted),
        "deducted_attributed": str(row.deducted_attributed),
        "deducted_allocated": str(row.deducted_allocated),
        "shortfall": str(row.shortfall),
        "excess": str(row.excess),
        "identified": row.party.identified,
        "provisional": row.provisional,
        "mapped_section": row.mapped_section,
        "pan_warning": row.party.pan_warning,
        "status": row.status,
        "status_label": tds_detail.status_label(row.status),
        "status_emoji": tds_detail.STATUS_EMOJI.get(row.status, ""),
        "in_totals": _counts(row),
        "money_at_stake": str(row.at_stake),
        "estimate": tds_detail.estimate(row.exposure),
        "applies": row.decision.applies,
    }


@dataclass
class Totals:
    parties: int = 0
    crossed: int = 0
    aggregate: Decimal = Decimal(0)
    tds: Decimal = Decimal(0)
    deducted: Decimal = Decimal(0)
    shortfall: Decimal = Decimal(0)
    excess: Decimal = Decimal(0)
    not_deducted: int = 0
    short: int = 0
    excess_parties: int = 0

    def add(self, row: PartySection) -> None:
        self.parties += 1
        self.aggregate += row.evaluation.aggregate
        if row.evaluation.crossed:
            self.crossed += 1
            self.tds += row.evaluation.tds
            self.deducted += row.deducted
            self.shortfall += row.shortfall
            self.excess += row.excess
        self.not_deducted += row.status == te.TdsStatus.NOT_DEDUCTED.value
        self.short += row.status == te.TdsStatus.SHORT.value
        self.excess_parties += row.status == te.TdsStatus.EXCESS.value

    def as_dict(self) -> dict:
        return {
            "parties": self.parties,
            "crossed": self.crossed,
            "aggregate": str(self.aggregate),
            "tds": str(self.tds),
            "deducted": str(self.deducted),
            "shortfall": str(self.shortfall),
            "excess": str(self.excess),
            "not_deducted": self.not_deducted,
            "short": self.short,
            "excess_parties": self.excess_parties,
        }


def build_report(analysis: ClientTds) -> dict:
    order = {key: i for i, key in enumerate(te.SECTION_KEYS)}
    sections: list[dict] = []
    grand = Totals()
    for key in sorted({r.section_key for r in analysis.rows}, key=lambda k: order.get(k, 99)):
        rows = [r for r in analysis.rows if r.section_key == key]
        rows.sort(key=lambda r: (-r.at_stake, not r.party.identified, -r.evaluation.aggregate))
        subtotal = Totals()
        for row in rows:
            if _counts(row):
                subtotal.add(row)
                grand.add(row)
        rule = rows[0].evaluation.rule
        pool = analysis.pools.get(key)
        unidentified = [r for r in rows if not r.party.identified]
        sections.append(
            {
                "key": key,
                "nature": rule.nature,
                "applies": analysis.payer.decision(key).applies,
                "reason": analysis.payer.decision(key).reason,
                "rows": [report_row(r) for r in rows],
                "subtotal": subtotal.as_dict(),
                "money_at_stake": str(sum((r.at_stake for r in rows), Decimal(0))),
                "pool": _pool(pool),
                "unidentified_amount": str(
                    sum((r.evaluation.aggregate for r in unidentified), Decimal(0))
                ),
            }
        )
    payer = analysis.payer
    return {
        "client_id": analysis.client_id,
        "client_name": analysis.client_name,
        "fy": analysis.fy,
        "has_data": analysis.has_data,
        "mapping_approved": analysis.mapping_approved,
        "payer": {
            "constitution": payer.constitution,
            "constitution_source": payer.constitution_source,
            "previous_fy": payer.previous_fy,
            "previous_turnover": (
                str(payer.previous_turnover) if payer.previous_turnover is not None else None
            ),
            "s194q": {"applies": payer.s194q.applies, "reason": payer.s194q.reason},
            "others": {
                "applies": payer.others.applies,
                "reason": payer.others.reason,
                "needs_confirmation": payer.others.needs_confirmation,
            },
            "audit_liable": payer.audit_liable,
            "audit_source": payer.audit_source,
            "rent_under_194ib": payer.rent_under_194ib,
            "s194ib": {"applies": payer.s194ib.applies, "reason": payer.s194ib.reason},
        },
        "sections": sorted(sections, key=lambda s: -Decimal(s["money_at_stake"])),
        "ranked": [
            report_row(r)
            for r in sorted(analysis.rows, key=lambda r: -r.at_stake)
            if r.at_stake > 0
        ],
        "money_at_stake": str(sum((r.at_stake for r in analysis.rows), Decimal(0))),
        "estimate_label": tc.ESTIMATE_LABEL,
        "return_fees": [
            {
                "return": f.label,
                "due": f.due.isoformat(),
                "tds": str(f.tds),
                "shortfall": str(f.shortfall),
                "days": f.days,
                "fee": str(f.fee),
                "deducted_in_books": str(f.deducted_in_books),
                "likely_unfiled": f.likely_unfiled,
            }
            for f in analysis.return_fees
        ],
        # at risk: periods with no TDS booked at all (the return was probably never filed)
        "return_fee_total": str(
            sum((f.fee for f in analysis.return_fees if f.likely_unfiled), Decimal(0))
        ),
        "return_fee_if_none_filed": str(sum((f.fee for f in analysis.return_fees), Decimal(0))),
        "pan_groups": _pan_groups(analysis),
        "totals": {
            **grand.as_dict(),
            "booked_without_party": str(
                sum((p.amount for p in analysis.pools.values()), Decimal(0))
            ),
            "unallocated": str(sum((p.unallocated for p in analysis.pools.values()), Decimal(0))),
            "unidentified_payments": str(
                sum(
                    (
                        r.evaluation.aggregate
                        for r in analysis.rows
                        if not r.party.identified and r.decision.applies is not False
                    ),
                    Decimal(0),
                )
            ),
        },
        "pending_choice": analysis.pending_choice,
        "data_quality": [
            {
                "party": r.party.name,
                "pan": r.party.pan,
                "problem": r.party.pan_warning,
                "level": te.pan_name_level(r.party.name, r.party.pan),
            }
            for r in analysis.rows
            if r.party.pan_warning
        ],
        "unmapped": [
            {"ledger": u.name, "kind": u.kind, "amount": str(u.amount), "lines": u.lines}
            for u in analysis.unmapped
        ],
    }


def _pan_groups(analysis: ClientTds) -> list[dict]:
    """One PAN behind several party ledgers in the same section: what Form 26Q will show
    for that deductee (the threshold is per payee, i.e. per PAN)."""
    groups: dict[tuple[str, str], list[PartySection]] = {}
    for row in analysis.rows:
        if row.party.identified and row.party.pan:
            groups.setdefault((row.party.pan, row.section_key), []).append(row)
    out = []
    for (pan, section), rows in sorted(groups.items()):
        if len({r.party.key for r in rows}) < 2:
            continue
        out.append(
            {
                "pan": pan,
                "section": section,
                "parties": [r.party.name for r in rows],
                "aggregate": str(sum((r.evaluation.aggregate for r in rows), Decimal(0))),
                "tds": str(sum((r.evaluation.tds for r in rows), Decimal(0))),
                "deducted": str(sum((r.deducted for r in rows), Decimal(0))),
            }
        )
    return out


def _pool(pool) -> dict | None:
    if pool is None:
        return None
    return {
        "amount": str(pool.amount),
        "attributed": str(pool.attributed),
        "attributed_to": pool.attributed_to,
        "allocated": str(pool.allocated),
        "unallocated": str(pool.unallocated),
    }


def summary_block(analysis: ClientTds) -> dict:
    """The TDS block of the Client Report and the TDS column of the Summary."""
    totals = Totals()
    for row in analysis.rows:
        if _counts(row):
            totals.add(row)
    return {
        "has_data": analysis.has_data,
        "mapping_approved": analysis.mapping_approved,
        "parties_crossed": totals.crossed,
        "tds_payable": str(totals.tds),
        "tds_deducted": str(totals.deducted),
        "not_deducted": str(totals.shortfall),
        "parties_not_deducted": totals.not_deducted,
        "parties_short": totals.short,
        "excess": str(totals.excess),
        "parties_excess": totals.excess_parties,
        "unidentified": sum(
            1 for r in analysis.rows if r.status == te.TdsStatus.UNIDENTIFIED.value
        ),
        "unallocated": str(sum((p.unallocated for p in analysis.pools.values()), Decimal(0))),
        "fy": analysis.fy,
        "unmapped_ledgers": len(analysis.unmapped),
        "s194q_applies": analysis.payer.s194q.applies,
    }


# ------------------------------------------------------------------- Excel

_MONEY = FMT_INDIAN_FULL
_BOLD = Font(bold=True)
_SUB_FILL = PatternFill("solid", fgColor="E8EEF5")
REPORT_HEADERS = [
    "Party / Payee name", "PAN", "Nature of payment", "Payee type", "PAN available?",
    "Single payment (largest)", "Aggregate payments in FY", "Threshold", "Threshold crossed?",
    "Crossed on (voucher)", "TDS base", "TDS rate applied", "TDS amount", "TDS deducted",
    "Shortfall", "Excess", "Status", "Money at stake",
]  # fmt: skip
_MONEY_COLS = (6, 7, 11, 13, 14, 15, 16, 18)


def _is_amount(value) -> bool:
    """True for a money string such as '1200.00' (not a rate text like '2%')."""
    if value is None:
        return False
    try:
        Decimal(str(value))
    except ArithmeticError:
        return False
    return True


def _money(ws, row: int, col: int, value) -> None:
    cell = ws.cell(row=row, column=col, value=float(Decimal(value)))
    cell.number_format = _MONEY
    cell.border = _BOX
    cell.alignment = Alignment(horizontal="right")


def report_xlsx(report: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "TDS Applicability"
    ws["A1"] = f"TDS APPLICABILITY - {report['client_name']} - FY {report['fy']}"
    ws["A1"].font = Font(bold=True, size=14)
    payer = report["payer"]
    ws["A2"] = f"194Q: {payer['s194q']['reason']}"
    ws["A3"] = f"194C / 194H / 194J / 194I: {payer['others']['reason']}"
    if not report["mapping_approved"]:
        ws["A4"] = "Ledger mapping not yet approved - figures are provisional."
        ws["A4"].font = Font(bold=True, color="B42318")
    widths = (34, 12, 40, 14, 10, 16, 18, 30, 18, 22, 16, 18, 14, 14, 14, 14, 34, 18)
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width
    row_no = 6
    _header_row(ws, row_no, REPORT_HEADERS)
    for section in report["sections"]:
        for item in section["rows"]:
            row_no += 1
            crossed = (
                f"{item['crossed_on'] or ''} ({item['crossed_voucher_no'] or ''})"
                if item["crossed_on"]
                else ""
            )
            values = [item["party"], item["pan"] or "not available", item["nature"], item["payee_type"],
                      item["pan_available"], None, None, item["threshold"], item["crossed"], crossed,
                      None, item["rate"], None, None, None, None, item["status_label"]]  # fmt: skip
            for col, value in enumerate(values, start=1):
                if col in _MONEY_COLS:
                    continue
                cell = ws.cell(row=row_no, column=col, value=value)
                cell.border = _BOX
            for col, key in zip(
                _MONEY_COLS,
                (
                    "single_payment",
                    "aggregate",
                    "base",
                    "tds",
                    "deducted",
                    "shortfall",
                    "excess",
                    "money_at_stake",
                ),
                strict=True,
            ):
                _money(ws, row_no, col, item[key])
            fill = SEVERITY_FILL.get(item["status"])
            if fill:
                ws.cell(row=row_no, column=17).fill = PatternFill("solid", fgColor=fill)
        row_no += 1
        sub = section["subtotal"]
        ws.cell(row=row_no, column=1, value=f"Subtotal {section['key']}").font = _BOLD
        for col, key in ((7, "aggregate"), (13, "tds"), (14, "deducted"), (15, "shortfall")):
            _money(ws, row_no, col, sub[key])
            ws.cell(row=row_no, column=col).font = _BOLD
        for col in range(1, 19):
            ws.cell(row=row_no, column=col).fill = _SUB_FILL
        pool = section["pool"]
        if pool:
            row_no += 1
            text = f"{section['key']}: TDS booked without a party"
            if pool["attributed_to"]:
                text += (
                    f" - attributed to {pool['attributed_to']} (the only party over the threshold)"
                )
            ws.cell(row=row_no, column=1, value=text)
            _money(ws, row_no, 14, pool["amount"])
            ws.cell(row=row_no, column=17, value=f"Unallocated: {pool['unallocated']}")
    row_no += 2
    totals = report["totals"]
    ws.cell(row=row_no, column=12, value="Total TDS").font = _BOLD
    _money(ws, row_no, 13, totals["tds"])
    _money(ws, row_no, 14, totals["deducted"])
    row_no += 1
    ws.cell(row=row_no, column=12, value="Total shortfall").font = _BOLD
    _money(ws, row_no, 15, totals["shortfall"])
    _money(ws, row_no, 16, totals["excess"])
    for label, key in (
        ("TDS booked without a party", "booked_without_party"),
        ("of which unallocated", "unallocated"),
        ("Payments with no party ledger (not in the totals)", "unidentified_payments"),
    ):
        row_no += 1
        ws.cell(row=row_no, column=12, value=label).font = _BOLD
        _money(ws, row_no, 14, totals[key])
    if report["unmapped"]:
        sheet = wb.create_sheet("Unmapped ledgers")
        _header_row(sheet, 1, ["Ledger", "Kind", "Net amount in FY", "Lines"])
        for i, item in enumerate(report["unmapped"], start=2):
            sheet.cell(row=i, column=1, value=item["ledger"])
            sheet.cell(row=i, column=2, value=item["kind"])
            _money(sheet, i, 3, item["amount"])
            sheet.cell(row=i, column=4, value=item["lines"])
        sheet.column_dimensions["A"].width = 44
        sheet.column_dimensions["C"].width = 20
    ws.freeze_panes = "B7"
    _ranked_sheet(wb, report)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _ranked_sheet(wb: Workbook, report: dict) -> None:
    """Every party/section with money at stake, largest first, with the estimates."""
    ws = wb.create_sheet("Ranked by money at stake", 0)
    ws["A1"] = f"TDS - ranked by money at stake - {report['client_name']} - FY {report['fy']}"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = report["estimate_label"]
    ws["A2"].font = Font(italic=True, color="B42318")
    head = ["Rank", "Party", "Section", "Status", "Aggregate", "TDS not deducted / indicative",
            "Interest 1%", "Interest 1.5% (alternative)", "Disallowance 30%",
            "Excess", "Money at stake"]  # fmt: skip
    _header_row(ws, 4, head)
    for i, item in enumerate(report["ranked"], start=1):
        r = 4 + i
        est = item["estimate"] or {}
        amounts = {
            line["item"].split(" (")[0].split(":")[0]: line["amount"]
            for line in est.get("lines", [])
        }
        values = [i, item["party"], item["section"], item["status_label"]]
        for col, value in enumerate(values, start=1):
            ws.cell(row=r, column=col, value=value).border = _BOX
        basis = next((v for k, v in amounts.items() if k.startswith("Basis")), "0")
        figures = [item["aggregate"], basis or "0",
                   amounts.get("Interest if not deducted") or "0",
                   amounts.get("Interest if deducted but not deposited") or "0",
                   amounts.get("Disallowance u/s 40(a)(ia)") or "0", item["excess"], item["money_at_stake"]]  # fmt: skip
        for col, value in enumerate(figures, start=5):
            _money(ws, r, col, value)
    ws.column_dimensions["B"].width = 40
    for col in "DEFGHIJK":
        ws.column_dimensions[col].width = 20
    r = 6 + len(report["ranked"])
    ws.cell(
        row=r, column=1, value="s.234E late fee - per RETURN (all parties), if not yet filed"
    ).font = Font(bold=True)
    _header_row(
        ws,
        r + 1,
        ["Return", "Due", "TDS in the return", "of which not deducted", "Days late",
         "Fee (capped)", "TDS booked in the period", "Fee counted"],
    )  # fmt: skip
    for i, fee in enumerate(report["return_fees"], start=r + 2):
        ws.cell(row=i, column=1, value=fee["return"])
        ws.cell(row=i, column=2, value=fee["due"])
        _money(ws, i, 3, fee["tds"])
        _money(ws, i, 4, fee["shortfall"])
        ws.cell(row=i, column=5, value=fee["days"])
        _money(ws, i, 6, fee["fee"])
        _money(ws, i, 7, fee["deducted_in_books"])
        counted = "at risk (no TDS booked)" if fee["likely_unfiled"] else "only if not filed"
        ws.cell(row=i, column=8, value=counted)


def party_xlsx(detail: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Working"
    party, section, threshold = detail["party"], detail["section"], detail["threshold"]
    lines = [
        (f"TDS WORKING - {detail['client_name']} - FY {detail['fy']}", None),
        (f"{detail['status_label']}", None),
        ("", None),
        ("Party", party["name"]),
        ("PAN", f"{party['pan'] or 'not available'} ({party['pan_source']})"),
        ("Payee type", f"{party['payee_type_label']} ({party['payee_source']})"),
        ("GSTIN", party["gstin"] or "-"),
        ("Section", f"{section['key']} - {section['nature']}"),
        ("Who must deduct", section["who_must_deduct"]),
        ("Why it applies", section["why"]["reason"]),
        ("Threshold", threshold["text"]),
        ("Crossed by", threshold["crossing_test"] or "not crossed"),
        (
            "Crossed on",
            f"{threshold['crossed_on'] or '-'} (voucher {threshold['crossed_voucher_no'] or '-'})",
        ),
    ]
    for i, (label, value) in enumerate(lines, start=1):
        ws.cell(row=i, column=1, value=label).font = _BOLD if i <= 2 or value else Font()
        if value is not None:
            ws.cell(row=i, column=2, value=value)
    row_no = len(lines) + 2
    _header_row(ws, row_no, ["Calculation", "Value", "Note"])
    for step in detail["calculation"]["steps"]:
        row_no += 1
        ws.cell(row=row_no, column=1, value=step["step"])
        if _is_amount(step["value"]):
            _money(ws, row_no, 2, step["value"])
        else:
            ws.cell(row=row_no, column=2, value=step["value"])
        ws.cell(row=row_no, column=3, value=step["note"])
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 90

    vouchers = wb.create_sheet("Vouchers")
    _header_row(vouchers, 1, ["Date", "Voucher type", "Voucher no", "Ledger", "Amount (excl. GST)",
                              "Running total", "Liable amount", "Rate %", "Crossing"])  # fmt: skip
    crossing_fill = PatternFill("solid", fgColor="FDE7E7")
    for i, line in enumerate(detail["vouchers"], start=2):
        vouchers.cell(row=i, column=1, value=date.fromisoformat(line["date"])).number_format = (
            "dd-mmm-yyyy"
        )
        vouchers.cell(row=i, column=2, value=line["voucher_type"])
        vouchers.cell(row=i, column=3, value=line["voucher_no"])
        vouchers.cell(row=i, column=4, value=line["ledger"])
        _money(vouchers, i, 5, line["amount"])
        _money(vouchers, i, 6, line["running_total"])
        _money(vouchers, i, 7, line["liable"])
        vouchers.cell(row=i, column=8, value=float(Decimal(line["rate"])))
        if line["is_crossing"]:
            vouchers.cell(row=i, column=9, value="Threshold crossed here")
            for col in range(1, 10):
                vouchers.cell(row=i, column=col).fill = crossing_fill
    for col, width in zip("ABCDEFGHI", (13, 16, 18, 40, 20, 20, 18, 8, 22), strict=True):
        vouchers.column_dimensions[col].width = width
    vouchers.auto_filter.ref = f"A1:I{max(1, len(detail['vouchers']) + 1)}"

    comp = wb.create_sheet("Compliance")
    comp["A1"] = detail["compliance"]["disclaimer"]
    comp["A1"].font = Font(bold=True, color="B42318")
    _header_row(comp, 3, ["Item", "Detail", "Figure"])
    for i, item in enumerate(detail["compliance"]["items"], start=4):
        comp.cell(row=i, column=1, value=item["item"])
        comp.cell(row=i, column=2, value=item["detail"]).alignment = Alignment(wrap_text=True)
        comp.cell(row=i, column=3, value=item["figure"])
    comp.column_dimensions["A"].width = 44
    comp.column_dimensions["B"].width = 100
    comp.column_dimensions["C"].width = 16
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------- PDF


def _inr(value) -> str:
    amount = Decimal(value)
    whole, _, paise = f"{abs(amount):.2f}".partition(".")
    head, tail = whole[:-3], whole[-3:]
    groups = []
    while head:
        groups.insert(0, head[-2:])
        head = head[:-2]
    text = ",".join([*groups, tail]) if groups else tail
    return f"{'-' if amount < 0 else ''}₹{text}.{paise}"


_GRID = [
    ("FONTNAME", (0, 0), (-1, -1), "DejaVu"),
    ("FONTSIZE", (0, 0), (-1, -1), 7),
    ("FONTNAME", (0, 0), (-1, 0), "DejaVu-Bold"),
    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BFC3C7")),
    ("VALIGN", (0, 0), (-1, -1), "TOP"),
]


def _doc(buffer, title: str) -> SimpleDocTemplate:
    return SimpleDocTemplate(
        buffer, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
        topMargin=12 * mm, bottomMargin=16 * mm, title=title, author="Turnover Analysis & Alert Tool",
    )  # fmt: skip


def report_pdf(report: dict) -> bytes:
    _register_fonts()
    styles = _styles()
    cell = styles["cell"].clone("tdscell", fontSize=7, leading=9)
    buffer = io.BytesIO()
    title = f"TDS APPLICABILITY – {report['client_name']} – FY {report['fy']}"
    doc = _doc(buffer, title)
    payer = report["payer"]
    story = [
        Paragraph(escape(title), styles["title"]),
        Paragraph(escape(f"194Q: {payer['s194q']['reason']}"), styles["cell"]),
        Paragraph(
            escape(f"194C / 194H / 194J / 194I: {payer['others']['reason']}"), styles["cell"]
        ),
    ]
    if not report["mapping_approved"]:
        story.append(
            Paragraph(
                "<b>Ledger mapping not yet approved – figures are provisional.</b>", styles["cell"]
            )
        )
    if report["ranked"]:
        story.append(
            Paragraph(
                "Ranked by money at stake (shortfall + 30% disallowance + interest at 1%)",
                styles["h2"],
            )
        )
        story.append(Paragraph(f"<b>{escape(report['estimate_label'])}</b>", styles["small"]))
        ranked = [["#", "Party", "Section", "Status", "Aggregate", "Money at stake"]]
        for i, item in enumerate(report["ranked"], start=1):
            ranked.append([str(i), Paragraph(escape(item["party"]), cell), item["section"],
                           Paragraph(escape(item["status_label"]), cell), _inr(item["aggregate"]),
                           _inr(item["money_at_stake"])])  # fmt: skip
        table = Table(ranked, colWidths=[w * mm for w in (8, 80, 18, 70, 35, 35)], repeatRows=1)
        table.setStyle(TableStyle([*_GRID, ("ALIGN", (4, 1), (5, -1), "RIGHT")]))
        story.append(table)
    story.append(Spacer(1, 6))
    head = ["Party", "PAN", "Section", "Payee", "Largest single", "Aggregate", "Crossed?", "Base",
            "Rate", "TDS", "Deducted", "Shortfall", "Status"]  # fmt: skip
    data = [head]
    style = list(_GRID)
    for section in report["sections"]:
        for item in section["rows"]:
            data.append([
                Paragraph(escape(item["party"]), cell), item["pan"] or "—", item["section"], item["payee_type"],
                _inr(item["single_payment"]), _inr(item["aggregate"]), "Yes" if item["crossed"].startswith("Yes") else "No",
                _inr(item["base"]), Paragraph(escape(item["rate"]), cell), _inr(item["tds"]), _inr(item["deducted"]),
                _inr(item["shortfall"]), Paragraph(escape(f"{item['status_emoji']} {item['status_label']}"), cell),
            ])  # fmt: skip
            fill = SEVERITY_FILL.get(item["status"])
            if fill:
                style.append(
                    (
                        "BACKGROUND",
                        (12, len(data) - 1),
                        (12, len(data) - 1),
                        colors.HexColor(f"#{fill}"),
                    )
                )
        sub = section["subtotal"]
        data.append([f"Subtotal {section['key']}", "", "", "", "", _inr(sub["aggregate"]), "", "", "",
                     _inr(sub["tds"]), _inr(sub["deducted"]), _inr(sub["shortfall"]), ""])  # fmt: skip
        style += [("BACKGROUND", (0, len(data) - 1), (-1, len(data) - 1), colors.HexColor("#E8EEF5")),
                  ("FONTNAME", (0, len(data) - 1), (-1, len(data) - 1), "DejaVu-Bold")]  # fmt: skip
    totals = report["totals"]
    data.append(["Total", "", "", "", "", _inr(totals["aggregate"]), "", "", "", _inr(totals["tds"]),
                 _inr(totals["deducted"]), _inr(totals["shortfall"]), ""])  # fmt: skip
    style += [("FONTNAME", (0, len(data) - 1), (-1, len(data) - 1), "DejaVu-Bold"),
              ("ALIGN", (4, 1), (11, -1), "RIGHT")]  # fmt: skip
    widths = [52, 20, 15, 18, 22, 25, 12, 25, 22, 18, 20, 20, 36]
    table = Table(data, colWidths=[w * mm for w in widths], repeatRows=1)
    table.setStyle(TableStyle(style))
    story.append(table)
    if report["unmapped"]:
        story.append(
            Paragraph("Unmapped ledgers with postings (not in the figures above)", styles["h2"])
        )
        story.append(
            Paragraph(
                escape(
                    "; ".join(f"{u['ledger']} ({_inr(u['amount'])})" for u in report["unmapped"])
                ),
                styles["cell"],
            )
        )
    doc.build(story, onFirstPage=_pdf_footer, onLaterPages=_pdf_footer)
    return buffer.getvalue()


def party_pdf(detail: dict) -> bytes:
    _register_fonts()
    styles = _styles()
    cell = styles["cell"].clone("tdscell2", fontSize=8, leading=10)
    buffer = io.BytesIO()
    party, section, threshold, calc = (
        detail["party"],
        detail["section"],
        detail["threshold"],
        detail["calculation"],
    )
    title = f"TDS WORKING – {party['name']} – {section['key']} – FY {detail['fy']}"
    doc = _doc(buffer, title)
    facts = [
        ["Client", detail["client_name"]],
        ["Status", f"{detail['status_emoji']} {detail['status_label']}"],
        ["Party / PAN", f"{party['name']} · PAN {party['pan'] or 'not available'} ({party['pan_source']})"],
        ["Payee type", f"{party['payee_type_label']} ({party['payee_source']})"],
        ["Section", f"{section['key']} – {section['nature']}"],
        ["Why it applies", section["why"]["reason"]],
        ["Threshold", threshold["text"]],
        ["Crossed", f"{threshold['crossing_test'] or 'not crossed'} on {threshold['crossed_on'] or '—'} (voucher {threshold['crossed_voucher_no'] or '—'})"],
    ]  # fmt: skip
    story = [Paragraph(escape(title), styles["title"])]
    table = Table(
        [[k, Paragraph(escape(str(v)), cell)] for k, v in facts], colWidths=[40 * mm, 225 * mm]
    )
    table.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), "DejaVu"), ("FONTSIZE", (0, 0), (-1, -1), 8),
                               ("FONTNAME", (0, 0), (0, -1), "DejaVu-Bold"), ("VALIGN", (0, 0), (-1, -1), "TOP")]))  # fmt: skip
    story += [table, Paragraph("Calculation", styles["h2"])]
    rows = [["Step", "Value", "Note"]]
    for step in calc["steps"]:
        value = step["value"]
        shown = _inr(value) if _is_amount(value) else (value or "")
        rows.append([step["step"], shown, Paragraph(escape(step["note"] or ""), cell)])
    steps = Table(rows, colWidths=[70 * mm, 40 * mm, 155 * mm], repeatRows=1)
    steps.setStyle(TableStyle([*_GRID, ("ALIGN", (1, 1), (1, -1), "RIGHT")]))
    story += [steps, Paragraph("Vouchers", styles["h2"])]
    rows = [
        ["Date", "Type", "No.", "Ledger", "Amount (excl. GST)", "Running total", "Liable", "Rate %"]
    ]
    style = list(_GRID)
    for line in detail["vouchers"]:
        rows.append([line["date"], line["voucher_type"], line["voucher_no"], Paragraph(escape(line["ledger"]), cell),
                     _inr(line["amount"]), _inr(line["running_total"]), _inr(line["liable"]), line["rate"]])  # fmt: skip
        if line["is_crossing"]:
            style.append(
                ("BACKGROUND", (0, len(rows) - 1), (-1, len(rows) - 1), colors.HexColor("#FDE7E7"))
            )
    vouchers = Table(
        rows, colWidths=[w * mm for w in (22, 25, 30, 75, 32, 32, 30, 15)], repeatRows=1
    )
    vouchers.setStyle(TableStyle([*style, ("ALIGN", (4, 1), (7, -1), "RIGHT")]))
    story += [vouchers, Paragraph("Compliance and consequences", styles["h2"]),
              Paragraph(f"<b>{escape(detail['compliance']['disclaimer'])}</b>", styles["small"])]  # fmt: skip
    rows = [["Item", "Detail", "Figure"]]
    for item in detail["compliance"]["items"]:
        figure = item["figure"]
        if _is_amount(figure):
            figure = _inr(figure)
        rows.append(
            [
                Paragraph(escape(item["item"]), cell),
                Paragraph(escape(item["detail"]), cell),
                figure or "",
            ]
        )
    comp = Table(rows, colWidths=[60 * mm, 175 * mm, 30 * mm], repeatRows=1)
    comp.setStyle(TableStyle(_GRID))
    story.append(comp)
    doc.build(story, onFirstPage=_pdf_footer, onLaterPages=_pdf_footer)
    return buffer.getvalue()
